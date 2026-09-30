"""Build an editable eight-slide REACTOR pitch deck from docs/submission/SLIDES.md."""

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def sections(markdown: str):
    blocks = re.split(r"(?=^## Slide \d+: )", markdown, flags=re.M)
    slides = []
    for block in blocks:
        match = re.match(r"## Slide (\d+): ([^\n]+)\n", block)
        if match:
            body = block[match.end():].split("\n## ", 1)[0].strip()
            body = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", body)
            body = body.replace("`", "").replace("**", "")
            slides.append((int(match.group(1)), match.group(2).strip(), body))
    if [number for number, _, _ in slides] != list(range(1, 9)):
        raise ValueError("Submission deck requires Slides 1–8 in order")
    return slides


def build(source: Path, output: Path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.33), Inches(7.5)
    for number, title, body in sections(source.read_text()):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(15, 23, 42)

        heading = slide.shapes.add_textbox(Inches(.9), Inches(.65), Inches(11.6), Inches(.9)).text_frame
        paragraph = heading.paragraphs[0]
        paragraph.text = f"{number:02d}  {title}"
        paragraph.font.size = Pt(34)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(125, 211, 252)

        text = slide.shapes.add_textbox(Inches(1), Inches(1.85), Inches(11.3), Inches(4.9)).text_frame
        text.word_wrap = True
        for index, line in enumerate(body.split("\n\n")):
            para = text.paragraphs[0] if index == 0 else text.add_paragraph()
            para.text = line.replace("\n", " ").strip()
            para.font.size = Pt(20 if number == 1 else 18)
            para.font.color.rgb = RGBColor(235, 241, 250)
            para.space_after = Pt(22)

        footer = slide.shapes.add_textbox(Inches(.9), Inches(7.05), Inches(11.4), Inches(.25)).text_frame
        footer.paragraphs[0].text = "REACTOR  |  Theme 05  |  Evidence and limitations are reported honestly"
        footer.paragraphs[0].font.size = Pt(10)
        footer.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)
        footer.paragraphs[0].alignment = PP_ALIGN.RIGHT
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path("docs/submission/SLIDES.md"))
    parser.add_argument("--output", type=Path, default=Path("docs/submission/REACTOR_Submission.pptx"))
    args = parser.parse_args()
    print(build(args.source, args.output))
